"""Auditoria temporária, somente leitura, do saldo do material 6020072.

Apenas imprime linhas deste código e contagens. Não modifica dados.
"""
from __future__ import annotations
from collections import Counter
import io
import pandas as pd
from openpyxl import load_workbook
from central_mrp_data import _download_source_cached, _download_derived_cached, api_call

TARGET = "06020072"
NAO_DISP = {"ALMOX. N.C.I", "ALMOX. N.C.R", "ASTEC", "G9-PROD", "G9-SEPA", "PROD. BARRAS", "PROD. CHAPAS", "PROD. PROC", "PROD. SUCATA", "PROD. TCTP", "QUALIDADE", "SETTA FIOS"}


def norm_code(value):
    text = str(value or "").strip()
    if text.endswith(".0"): text = text[:-2]
    if not text.isdigit(): return ""
    return text.zfill(8)


def number(value):
    if isinstance(value, (float, int)):
        return float(value)
    return pd.to_numeric(str(value or "0").strip().replace(".", "").replace(",", "."), errors="coerce")


def metadata(kind, key):
    action = "derived_state" if kind == "derived" else "source_state"
    # O download independe das versões: cache keys só protegem recargas.
    return {}


def load_raw(key):
    blob = _download_source_cached(key, 0, "")
    wb = load_workbook(io.BytesIO(blob), read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        print("RAW", key, "sheet", ws.title, "columns", ws.max_column, "rows", ws.max_row)
        selected = []
        for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if row_idx <= 2: continue
            if not row: continue
            if norm_code(row[0]) == TARGET:
                selected.append((row_idx, row))
        return selected
    finally:
        wb.close()


def print_rows(label, rows, is_address=False):
    values = []
    print(label, "matches", len(rows))
    for row_num, row in rows:
        addr = str(row[3] or "").strip() if is_address else ""
        qty = row[7] if len(row) > 7 else None
        numeric = number(qty)
        values.append(numeric if pd.notna(numeric) else 0.0)
        print("LINE", row_num, "CODE", norm_code(row[0]), "DESCRIPTION_OR_ADDRESS", str(row[3])[:80], "RAW_H_TYPE", type(qty).__name__, "RAW_H", str(qty)[:60], "PARSED_H", numeric, "NON_AVAILABLE", addr in NAO_DISP if is_address else "")
    print("TOTAL",label,sum(values))
    if is_address:
        print("NON_AVAILABLE_TOTAL",sum(number(row[7]) or 0 for _, row in rows if str(row[3] or "").strip() in NAO_DISP))


def audit_derived(key):
    frame = _download_derived_cached(key, "audit", "audit")
    print("DERIVED", key, "shape", frame.shape, "columns", list(frame.columns))
    code_field = next((c for c in frame.columns if c.upper() in {"COD_MATERIAL","CÓDIGO","CODIGO"}), None)
    if not code_field:
        print("DERIVED_MISSING_CODE_FIELD")
        return
    filt = frame[frame[code_field].map(norm_code).eq(TARGET)]
    print("DERIVED_MATCHES", key, len(filt))
    if len(filt):
        print("DERIVED_ROWS", filt.head(15).to_json(orient="records",force_ascii=False))


try:
    audit_derived("estoque_tratado")
    audit_derived("relatorio_mrp")
    print_rows("ANALITICO", load_raw("analitico"))
    print_rows("ENDERECO", load_raw("endereco"), True)
except Exception as exc:
    print("AUDIT_FAILED", type(exc).__name__, str(exc)[:300])
    raise
